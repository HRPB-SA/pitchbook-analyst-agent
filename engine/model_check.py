"""Independent arithmetic check on a generated model.

LibreOffice cannot evaluate spreadsheets in every environment, and a model
nobody has computed is a model nobody should trust. This module reads the
inputs back out of a built workbook, recomputes the whole chain in Python
following the same rules the formulas encode, and reports what the model will
produce when Excel opens it.

It is a second implementation, not a re-reading of the first: if the formulas
and this code disagree about what the model means, that disagreement is the
finding.

    python3 -m engine modelcheck <workbook.xlsx>
"""
from __future__ import annotations

import datetime as _dt

import openpyxl


def _label_map(ws, col="B", val_col="D", limit=80):
    """Read an Inputs sheet as {label: value}."""
    out = {}
    for r in range(1, limit):
        k = ws[f"{col}{r}"].value
        v = ws[f"{val_col}{r}"].value
        if isinstance(k, str) and v is not None and not isinstance(v, str):
            out[k.strip()] = v
    return out


def check(path: str, mark_bn: float = None) -> dict:
    wb = openpyxl.load_workbook(path)
    inp = _label_map(wb["Inputs"])
    this_year = _dt.date.today().year

    def g(*names, default=None):
        for n in names:
            for k, v in inp.items():
                if k.lower().startswith(n.lower()):
                    return v
        return default

    rev0 = g("Latest annualised revenue")
    growth0 = g("Latest revenue growth")
    decay = g("Growth decay")
    tg = g("Terminal growth")
    scen = 1.0
    gm0 = g("Gross margin (%)")
    gm_drift = g("Gross margin drift")
    sm = g("Sales and marketing")
    rd = g("Research and development")
    ga = g("General and administrative")
    lev = g("Operating expense leverage")
    wacc = g("Discount rate")
    tax = g("Tax rate")
    exit_x = g("Exit multiple")
    capex_pct = 0.04
    for r_ in range(1, 200):
        if str(wb["Model"][f"B{r_}"].value or "").startswith("Capital expenditure and working capital (%"):
            v_ = wb["Model"][f"D{r_}"].value
            if isinstance(v_, (int, float)):
                capex_pct = float(v_)
            break
    val = g("Latest valuation")
    eq = g("Equity raised")

    missing = [n for n, v in (("revenue", rev0), ("growth", growth0), ("wacc", wacc),
                              ("terminal growth", tg)) if v is None]
    if missing:
        return {"ok": False, "error": f"inputs missing: {', '.join(missing)}"}

    # ---- the same chain the formulas encode ---------------------------------
    # The label "Revenue" appears twice: once as a section banner with no data,
    # once on the row that carries it. Take the row that actually has numbers.
    ms = wb["Model"]
    hist = {}
    for r in range(1, 200):
        if ms[f"B{r}"].value != "Revenue":
            continue
        found = {}
        for c in range(4, 20):
            cell = ms.cell(row=r, column=c)
            if isinstance(cell.value, (int, float)):
                found[2022 + (c - 4)] = float(cell.value)
        if found:
            hist = found
            break
    if not hist:
        return {"ok": False, "error": "no historical revenue row found in the Model sheet"}

    rows, rev, gm = [], None, gm0
    for yr in range(2022, 2033):
        if yr <= this_year:
            rev = hist.get(yr, rev)
            growth = None
            gm_y = gm0
        else:
            growth = max(tg, (rows[-1]["growth"] if rows[-1]["growth"] is not None
                              else growth0) - decay) * scen
            rev = rows[-1]["revenue"] * (1 + growth)
            gm_y = rows[-1]["gm"] + gm_drift
        if rev is None:
            continue
        gp = rev * gm_y
        step = max(0, yr - this_year)
        opex = sum(rev * max(0.02, base + lev * step) for base in (sm, rd, ga))
        ebit = gp - opex
        t = -max(0.0, ebit) * tax
        capex = -rev * capex_pct
        fcf = ebit + t + capex
        disc = 0.0 if yr <= this_year else 1 / (1 + wacc) ** (yr - this_year)
        rows.append({"year": yr, "growth": growth, "revenue": rev, "gm": gm_y,
                     "gross_profit": gp, "opex": opex, "ebit": ebit,
                     "margin": ebit / rev if rev else 0, "fcf": fcf,
                     "disc": disc, "pv": fcf * disc})

    fwd = [r for r in rows if r["year"] > this_year]
    sum_pv = sum(r["pv"] for r in fwd)
    last = rows[-1]
    tv_perp = last["fcf"] * (1 + tg) / (wacc - tg) * last["disc"]
    tv_exit = last["revenue"] * exit_x * last["disc"]
    # perpetuity is the base; the multiple is a cross-check, not an average
    tv = tv_perp
    spread = (max(tv_perp, tv_exit) / min(tv_perp, tv_exit)) if min(tv_perp, tv_exit) else 0
    ev = sum_pv + tv

    # ---- sanity gates -------------------------------------------------------
    checks = []

    def gate(name, ok, detail, severity="error"):
        """severity="error" is a mechanical fault the model must not ship with.
        severity="caution" is an analytical characteristic worth knowing: the
        model is working, and what it is telling you deserves attention."""
        checks.append({"check": name, "pass": bool(ok), "detail": detail,
                       "severity": severity})

    gate("Revenue rises every forecast year",
         all(fwd[i]["revenue"] > fwd[i - 1]["revenue"] for i in range(1, len(fwd))),
         f"{fwd[0]['revenue']:,.0f} -> {last['revenue']:,.0f} ($M)")
    gate("Growth decays toward the terminal rate",
         abs(last["growth"] - tg) < 0.2 or last["growth"] <= fwd[0]["growth"],
         f"{fwd[0]['growth']:.1%} -> {last['growth']:.1%} (terminal {tg:.1%})")
    gate("Gross margin stays between 20% and 100%",
         all(0.2 < r["gm"] < 1.0 for r in rows),
         f"{rows[0]['gm']:.1%} -> {last['gm']:.1%}")
    gate("Gross profit never exceeds revenue",
         all(r["gross_profit"] <= r["revenue"] + 1e-6 for r in rows), "by construction")
    gate("Discount rate exceeds terminal growth",
         wacc > tg, f"{wacc:.1%} vs {tg:.1%}")
    gate("Enterprise value is positive", ev > 0, f"${ev:,.0f}M")
    gate("Operating margin reaches a plausible level",
         -1.0 < last["margin"] < 0.6, f"{last['margin']:.1%} by {last['year']}")
    gate("Terminal value is under 75% of enterprise value",
         (tv / ev < 0.75) if ev else False,
         (f"{tv / ev:.0%} of EV — most of the value sits beyond the forecast, so the "
          f"answer rests on the terminal assumption rather than the projection")
         if ev else "n/a", "caution")
    gate("Terminal methods agree within 1.5x", spread and spread <= 1.5,
         f"perpetuity ${tv_perp:,.0f}M vs multiple ${tv_exit:,.0f}M ({spread:.1f}x apart) "
         f"— averaging them would hide the disagreement", "caution")

    mark = (mark_bn * 1000) if mark_bn else (val if val else None)
    out = {
        "ok": all(c["pass"] for c in checks if c["severity"] == "error"),
        "cautions": [c for c in checks if not c["pass"] and c["severity"] == "caution"],
        "path": path,
        "rows": rows,
        "sum_pv": sum_pv, "tv_perp": tv_perp, "tv_exit": tv_exit,
        "terminal_value": tv, "tv_spread": spread, "enterprise_value": ev,
        "mark": mark,
        "ev_vs_mark": (ev / mark) if mark else None,
        "mark_vs_revenue": (mark / hist.get(this_year)) if mark and hist.get(this_year) else None,
        "mark_per_equity": (mark / eq) if mark and eq else None,
        "checks": checks,
    }
    return out


def report(res: dict) -> str:
    if not res.get("ok") and res.get("error"):
        return f"FAILED: {res['error']}"
    L = [f"Model arithmetic check — {res['path']}", ""]
    L.append(f"{'Year':<7}{'Growth':>9}{'Revenue':>12}{'GM':>8}{'Op profit':>12}"
             f"{'Op margin':>11}{'Free CF':>12}")
    for r in res["rows"]:
        gtxt = f"{r['growth']:.1%}" if r["growth"] is not None else "—"
        L.append(f"{r['year']:<7}{gtxt:>9}{r['revenue']:>12,.0f}{r['gm']:>8.1%}"
                 f"{r['ebit']:>12,.0f}{r['margin']:>11.1%}{r['fcf']:>12,.0f}")
    L += ["", f"Sum of discounted cash flows   ${res['sum_pv']:>14,.0f}M",
          f"Terminal value (perpetuity)    ${res['tv_perp']:>14,.0f}M",
          f"Terminal value (exit multiple) ${res['tv_exit']:>14,.0f}M",
          f"Terminal value used (perpetuity)${res['terminal_value']:>14,.0f}M",
          f"Spread between the two methods  {res['tv_spread']:>14.1f}x",
          f"ENTERPRISE VALUE               ${res['enterprise_value']:>14,.0f}M"]
    if res.get("mark"):
        L.append(f"Stored mark                    ${res['mark']:>14,.0f}M")
        L.append(f"Model against the mark          {res['ev_vs_mark']:>14.2f}x")
        if res.get("mark_vs_revenue"):
            L.append(f"Mark against revenue            {res['mark_vs_revenue']:>14.1f}x")
        if res.get("mark_per_equity"):
            L.append(f"Mark per dollar of equity       {res['mark_per_equity']:>14.1f}x")
    L += ["", "Mechanical gates (must pass):"]
    for c in res["checks"]:
        if c["severity"] != "error":
            continue
        L.append(f"  [{'PASS' if c['pass'] else 'FAIL'}] {c['check']:<48} {c['detail']}")
    cautions = [c for c in res["checks"] if c["severity"] == "caution"]
    if cautions:
        L += ["", "Analytical cautions (the model works; read what it is saying):"]
        for c in cautions:
            L.append(f"  [{'ok' if c['pass'] else 'FLAG'}] {c['check']:<48} {c['detail']}")
    L.append("")
    L.append("MECHANICALLY SOUND" if res["ok"] else "MECHANICAL FAULT — do not ship")
    if res.get("mark") and res.get("ev_vs_mark"):
        L.append(f"The model values this business at {res['ev_vs_mark']:.2f}x the price "
                 f"last paid for it. That gap is the analysis, not an error.")
    return "\n".join(L)
