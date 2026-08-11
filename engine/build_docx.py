"""Template-driven report builder. Generalizes the Databricks 'Brick by Brick'
builder: US Letter, Georgia body, navy headings, pageBreakBefore H1s, dual-width
DXA tables with no-split rows and repeated headers, 300-DPI embedded charts,
page-x-of-y footer with the embargo line, cover with hero tiles + snapshot
table, measured Contents page. Em-dash scan fails the build.

Input: a report directory containing
    report.json      meta (see below)
    blocks/*.json    ordered section blocks (engine.compose grammar)
    charts/*.png     report-local charts (falls back to repo charts/)
Output: output/<outname>.docx (+ .pdf when soffice is available)

report.json:
{
  "outname": "Acme_Update_Aug2026",
  "kicker": "COMPANY UPDATE", "title": "Acme, Inc.",
  "subtitle": "...", "dek": "...",
  "header_line": "ACME, INC.  |  COMPANY UPDATE  |  AUGUST 10, 2026",
  "footer_note": "Quality-valuation coefficient embargoed; ... | Not investment advice",
  "confidential_line": "CONFIDENTIAL  |  For institutional recipients only  |  Not for redistribution",
  "hero_tiles": [{"label": ..., "value": ..., "caption": ..., "accent": true}, x3],
  "snapshot": {table spec}, "toc": true,
  "exhibits": [[1, "caption"], ...],
  "byline": "Prepared by ...  |  August 10, 2026",
  "disclosure": "..."
}
"""
from __future__ import annotations
import json, os, glob, subprocess, shutil
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from . import compose, qa

NAVY = RGBColor(0x1F, 0x2A, 0x44); SLATE = RGBColor(0x35, 0x50, 0x6E)
GREEN = RGBColor(0x1C, 0x5D, 0x46); GREY = RGBColor(0x6B, 0x72, 0x80)
INK = RGBColor(0x26, 0x2B, 0x35); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PB_BLUE = RGBColor(0x2F, 0x54, 0x96); PB_LINK = RGBColor(0x05, 0x63, 0xC1)
SOFT_HEX, NAVY_HEX, BORDER_HEX, SUM_HEX = "F2F5F9", "1F2A44", "C9D2DE", "E2E8F0"
CONTENT_W = 9360
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Report themes. "house" is the Georgia/navy institutional deep-dive look;
# "analyst_note" matches report_template/Vertical_Analyst_Note_10.docx
# (Segoe UI, blue Heading 2 sections, production metadata sheet as page 1,
# numbered References endnotes, no header/footer chrome).
THEMES = {
    "house": {
        "font": "Georgia", "body_size": 10.5, "body_color": INK,
        "h1": (16, NAVY, True), "h2": (12.5, NAVY, False), "h3": (11, SLATE, False),
        "accent": NAVY, "accent_hex": NAVY_HEX, "muted": GREY,
        "cover": "institutional", "chrome": True, "caption_max_words": None,
    },
    "analyst_note": {
        "font": "Segoe UI", "body_size": 10, "body_color": RGBColor(0x1A, 0x1A, 0x1A),
        "h1": (14, PB_BLUE, True), "h2": (12, PB_BLUE, False), "h3": (10.5, PB_BLUE, False),
        "accent": PB_BLUE, "accent_hex": "2F5496", "muted": RGBColor(0x60, 0x5E, 0x5C),
        "cover": "meta_sheet", "chrome": False, "caption_max_words": 30,
    },
}


class ReportBuilder:
    def __init__(self, report_dir):
        self.dir = os.path.abspath(report_dir)
        with open(os.path.join(self.dir, "report.json"), encoding="utf-8") as fh:
            self.meta = json.load(fh)
        self.T = THEMES[self.meta.get("theme", "house")]
        self.charts_dirs = [os.path.join(self.dir, "charts"),
                            os.path.join(REPO, "charts")]
        self.outdir = os.path.join(self.dir, "output")
        os.makedirs(self.outdir, exist_ok=True)
        self.doc = Document()
        self._setup_page()
        self._setup_styles()

    # ------------------------------------------------------------- setup
    def _setup_page(self):
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Emu(12240 * 635), Emu(15840 * 635)
        for a in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(sec, a, Inches(1))
        sec.different_first_page_header_footer = True
        self.sec = sec

    def _setup_styles(self):
        T = self.T
        st = self.doc.styles["Normal"]
        st.font.name = T["font"]; st.font.size = Pt(T["body_size"])
        st.font.color.rgb = T["body_color"]
        st.paragraph_format.space_after = Pt(7)
        st.paragraph_format.line_spacing = 1.12
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for attr in ("w:ascii", "w:hAnsi"):
            rf.set(qn(attr), T["font"])
        h1s, h1c, h1pb = T["h1"]; h2s, h2c, _ = T["h2"]; h3s, h3c, _ = T["h3"]
        self._style_heading("Heading 1", h1s, h1c, page_break=h1pb, before=0, after=10)
        self._style_heading("Heading 2", h2s, h2c, before=14, after=5)
        self._style_heading("Heading 3", h3s, h3c, before=10, after=4,
                            italic=(self.meta.get("theme", "house") == "house"))

    def _style_heading(self, name, size, color, bold=True, before=10, after=6,
                       page_break=False, italic=False):
        s = self.doc.styles[name]
        s.font.name = self.T["font"]; s.font.size = Pt(size)
        s.font.bold = bold; s.font.italic = italic; s.font.color.rgb = color
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.page_break_before = page_break
        s.paragraph_format.keep_with_next = True
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for attr in ("w:ascii", "w:hAnsi"):
            rf.set(qn(attr), self.T["font"])

    # -------------------------------------------------------- primitives
    @staticmethod
    def _shade(cell, hexfill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hexfill)
        tcPr.append(shd)

    @staticmethod
    def _cell_w(cell, dxa):
        tcPr = cell._tc.get_or_add_tcPr()
        tcW = tcPr.find(qn("w:tcW"))
        if tcW is None:
            tcW = OxmlElement("w:tcW"); tcPr.append(tcW)
        tcW.set(qn("w:w"), str(dxa)); tcW.set(qn("w:type"), "dxa")

    @staticmethod
    def _table_w(table, dxa):
        tblPr = table._tbl.tblPr
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
        tblW.set(qn("w:w"), str(dxa)); tblW.set(qn("w:type"), "dxa")

    @staticmethod
    def _borders(table, hexcolor=BORDER_HEX, sz="4"):
        tblPr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), sz)
            e.set(qn("w:space"), "0"); e.set(qn("w:color"), hexcolor)
            borders.append(e)
        tblPr.append(borders)

    @staticmethod
    def _cell_margins(table, top=40, bottom=40, left=80, right=80):
        tblPr = table._tbl.tblPr
        m = OxmlElement("w:tblCellMar")
        for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            e = OxmlElement(f"w:{name}")
            e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa")
            m.append(e)
        tblPr.append(m)

    @staticmethod
    def _no_split(row, header=False):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if header:
            th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true")
            trPr.append(th)

    def _cell_text(self, cell, text, size=8.6, bold=False, color=INK, align="L", italic=False):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = {"L": WD_ALIGN_PARAGRAPH.LEFT, "R": WD_ALIGN_PARAGRAPH.RIGHT,
                       "C": WD_ALIGN_PARAGRAPH.CENTER}[align]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.font.name = self.T["font"]; r.font.size = Pt(size)
        r.font.bold = bold; r.font.italic = italic; r.font.color.rgb = color

    def _field(self, paragraph, instr):
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), instr)
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"
        r.append(t); fld.append(r)
        paragraph._p.append(fld)

    def para(self, text, size=10.5, color=INK, italic=False, bold=False,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=7):
        p = self.doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        r = p.add_run(text)
        r.font.size = Pt(size); r.font.color.rgb = color
        r.font.italic = italic; r.font.bold = bold
        return p

    def bullets(self, items):
        for it in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.28)
            r = p.add_run(it)
            r.font.size = Pt(10.5); r.font.color.rgb = INK

    @staticmethod
    def _col_widths(header, rows):
        n = len(header)
        weights = []
        for j in range(n):
            longest = max([len(str(header[j]))] + [len(str(r[j])) for r in rows])
            weights.append(min(max(longest, 6), 60))
        total = sum(weights)
        return [max(int(CONTENT_W * w / total), 700) for w in weights]

    def table_block(self, spec):
        header, rows = spec["header"], spec["rows"]
        align = spec.get("align", "L" * len(header))
        if spec.get("title"):
            tp = self.para(spec["title"], size=9.5, color=NAVY, bold=True,
                           align=WD_ALIGN_PARAGRAPH.LEFT, space_after=3)
            tp.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=len(rows) + 1, cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        self._table_w(t, CONTENT_W); self._borders(t); self._cell_margins(t)
        widths = spec.get("widths") or self._col_widths(header, rows)
        self._no_split(t.rows[0], header=True)
        for r_ in t.rows[1:]:
            self._no_split(r_)
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]
            self._cell_w(c, widths[j]); self._shade(c, self.T["accent_hex"])
            self._cell_text(c, str(h), bold=True, color=WHITE, align=align[j])
        for i, row in enumerate(rows):
            last = (i == len(rows) - 1)
            summary = last and any("total" in str(x).lower() or "Composite" in str(x)
                                   for x in row)
            for j, val in enumerate(row):
                c = t.rows[i + 1].cells[j]
                self._cell_w(c, widths[j])
                if summary:
                    self._shade(c, SUM_HEX)
                elif i % 2 == 1:
                    self._shade(c, SOFT_HEX)
                self._cell_text(c, str(val), size=8.4, align=align[j],
                                bold=summary)
        if spec.get("source"):
            self.para(spec["source"], size=7.8, color=GREY, italic=True,
                      align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10)
        else:
            self.para("", size=4, space_after=4)

    def fig(self, fname, caption):
        path = next((os.path.join(d, fname) for d in self.charts_dirs
                     if os.path.exists(os.path.join(d, fname))), None)
        if path is None:
            raise FileNotFoundError(f"chart not found in {self.charts_dirs}: {fname}")
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Inches(6.5))
        self.para(caption, size=7.8, color=GREY, italic=True,
                  align=WD_ALIGN_PARAGRAPH.LEFT, space_after=12)

    def render_blocks(self, blocks):
        for b in blocks:
            kind = b[0]
            if kind in ("h1", "h2", "h3"):
                self.doc.add_heading(b[1], level=int(kind[1]))
            elif kind == "p":
                self.para(b[1])
            elif kind == "bullets":
                self.bullets(b[1])
            elif kind == "table":
                self.table_block(b[1])
            elif kind == "fig":
                self.fig(b[1], b[2])
            else:
                raise ValueError(kind)

    # ------------------------------------------------------ furniture
    def header_footer(self):
        m = self.meta
        hp = self.sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run(m.get("header_line", ""))
        hr.font.name = self.T["font"]; hr.font.size = Pt(7.5)
        hr.font.color.rgb = GREY; hr.font.bold = True
        pPr = hp._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "4"); bottom.set(qn("w:color"), BORDER_HEX)
        pbdr.append(bottom); pPr.append(pbdr)

        fp = self.sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f1 = fp.add_run(m.get("footer_note", "Not investment advice"))
        f1.font.name = self.T["font"]; f1.font.size = Pt(6.8)
        f1.font.color.rgb = GREY; f1.font.italic = True
        fp2 = self.sec.footer.add_paragraph()
        fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for txt, fld in (("Page ", "PAGE"), (" of ", "NUMPAGES")):
            r = fp2.add_run(txt)
            r.font.name = self.T["font"]; r.font.size = Pt(7.5); r.font.color.rgb = GREY
            self._field(fp2, fld)
        for fldr in fp2._p.findall(qn("w:fldSimple")):
            for rr in fldr.findall(qn("w:r")):
                rpr = OxmlElement("w:rPr")
                fonts = OxmlElement("w:rFonts")
                fonts.set(qn("w:ascii"), self.T["font"]); fonts.set(qn("w:hAnsi"), self.T["font"])
                szel = OxmlElement("w:sz"); szel.set(qn("w:val"), "15")
                colel = OxmlElement("w:color"); colel.set(qn("w:val"), "6B7280")
                rpr.append(fonts); rpr.append(szel); rpr.append(colel)
                rr.insert(0, rpr)
        fpf = self.sec.first_page_footer.paragraphs[0]
        fpf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fpf.add_run(m.get("confidential_line",
                              "CONFIDENTIAL  |  For institutional recipients only  |  Not for redistribution"))
        r.font.name = self.T["font"]; r.font.size = Pt(7.5)
        r.font.color.rgb = GREY; r.font.italic = True

    def cover(self):
        m = self.meta
        self.para("", size=10, space_after=26)
        self.para(m.get("kicker", "RESEARCH NOTE"), size=10, color=GREY, bold=True,
                  align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(m["title"])
        r.font.name = self.T["font"]; r.font.size = Pt(30)
        r.font.bold = True; r.font.color.rgb = NAVY
        if m.get("subtitle"):
            self.para(m["subtitle"], size=14, color=SLATE, italic=True,
                      align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
        if m.get("dek"):
            self.para(m["dek"], size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=14)
        tiles = m.get("hero_tiles") or []
        if tiles:
            t = self.doc.add_table(rows=1, cols=len(tiles))
            self._table_w(t, CONTENT_W)
            self._borders(t, hexcolor=NAVY_HEX, sz="8")
            self._cell_margins(t, top=130, bottom=130, left=120, right=120)
            for j, tile in enumerate(tiles):
                c = t.rows[0].cells[j]
                self._cell_w(c, CONTENT_W // len(tiles))
                self._shade(c, "EAF2EE" if tile.get("accent") else SOFT_HEX)
                c.text = ""
                p1 = c.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p1.paragraph_format.space_after = Pt(2)
                r1 = p1.add_run(tile["label"])
                r1.font.name = self.T["font"]; r1.font.size = Pt(7.5)
                r1.font.bold = True; r1.font.color.rgb = GREY
                p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p2.paragraph_format.space_after = Pt(2)
                r2 = p2.add_run(tile["value"])
                r2.font.name = self.T["font"]; r2.font.size = Pt(13); r2.font.bold = True
                r2.font.color.rgb = GREEN if tile.get("accent") else NAVY
                p3 = c.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p3.paragraph_format.space_after = Pt(0)
                r3 = p3.add_run(tile.get("caption", ""))
                r3.font.name = self.T["font"]; r3.font.size = Pt(7.3); r3.font.color.rgb = GREY
            self.para("", size=6, space_after=8)
        if m.get("snapshot"):
            self.table_block(m["snapshot"])
        if m.get("byline"):
            self.para(m["byline"], size=8.2, color=GREY, italic=True,
                      align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)

    def contents(self, section_titles):
        m = self.meta
        pages = {}
        pfile = os.path.join(self.dir, "toc_pages.json")
        if os.path.exists(pfile):
            with open(pfile, encoding="utf-8") as fh:
                pages = json.load(fh)
        self.doc.add_heading("Contents", level=1)
        t = self.doc.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        self._table_w(t, CONTENT_W)
        self._cell_margins(t, top=26, bottom=26)
        for title in section_titles:
            r_ = t.add_row(); self._no_split(r_)
            self._cell_w(r_.cells[0], CONTENT_W - 900)
            self._cell_w(r_.cells[1], 900)
            self._cell_text(r_.cells[0], title, size=9.6)
            self._cell_text(r_.cells[1], str(pages.get(title, "")), size=9.6, align="R", color=GREY)
        ex = m.get("exhibits") or []
        if ex:
            self.para("", size=6, space_after=6)
            self.para("Exhibits", size=11, color=NAVY, bold=True,
                      align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
            for num, cap in ex:
                self.para(f"Figure {num}.  {cap}", size=8.6, color=INK,
                          align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)

    def _kv(self, label, value, bold_label=True):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label + ": " if label else "")
        r.font.size = Pt(10); r.font.bold = bold_label
        r2 = p.add_run(str(value))
        r2.font.size = Pt(10)
        return p

    def cover_meta_sheet(self, section_titles):
        """Page 1 of the Vertical Analyst Note: the production metadata sheet
        (report_template/Vertical_Analyst_Note_10.docx, page 1), filled."""
        m, prod = self.meta, self.meta.get("production", {})
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(m["title"])
        r.font.size = Pt(14); r.font.bold = True
        if m.get("subtitle"):
            p2 = self.doc.add_paragraph()
            p2.paragraph_format.space_after = Pt(8)
            r2 = p2.add_run(m["subtitle"])
            r2.font.size = Pt(11); r2.font.bold = True
        self._kv("Data filepath", prod.get("data_filepath", ""))
        self._kv("Chart as-of date", prod.get("chart_as_of", ""))
        self._kv("Chart geography", prod.get("chart_geography", ""))
        self._kv("Research type (Core/Quant, Emerging Tech, Industry)",
                 prod.get("research_type", ""))
        self._kv("Access level", prod.get("access_level", ""))
        self.para("", size=4, space_after=2)
        self._kv("Chart count", prod.get("chart_count", ""))
        self._kv("Table count", prod.get("table_count", ""))
        self.doc.add_heading("Credits", level=2)
        for c in prod.get("credits", []):
            self._kv("", f"{c['name']}  {c['title']}", bold_label=False)
        if prod.get("published"):
            self.para(f"Published on {prod['published']}", size=10, space_after=8,
                      align=WD_ALIGN_PARAGRAPH.LEFT)
        self.doc.add_heading("Contents", level=2)
        for t in section_titles:
            self.para(t, size=10, space_after=1, align=WD_ALIGN_PARAGRAPH.LEFT)
        if self.meta.get("references"):
            self.para("References", size=10, space_after=6,
                      align=WD_ALIGN_PARAGRAPH.LEFT)
        lp = prod.get("landing_page")
        if lp:
            self.doc.add_heading("Landing page", level=2)
            self.para("CHART: " + lp.get("chart", ""), size=10,
                      color=self.T["accent"], italic=True,
                      align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
            self.para(lp.get("caption", ""), size=10,
                      align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)
            self.para(lp.get("header", ""), size=10, bold=True,
                      align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
            self.para(lp.get("body", ""), size=10,
                      align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8)
        picks = prod.get("report_picks")
        if picks:
            self.doc.add_heading("Report picks", level=2)
            for pick in picks:
                self.para(pick, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=1)

    def references_section(self):
        refs = self.meta.get("references") or []
        if not refs:
            return
        self.doc.add_heading("References", level=2)
        for i, ref in enumerate(refs, start=1):
            text = ref.get("text", "") if isinstance(ref, dict) else str(ref)
            url = ref.get("url", "") if isinstance(ref, dict) else ""
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            rn = p.add_run(f"{i}. ")
            rn.font.size = Pt(9); rn.font.bold = True
            rt = p.add_run(text + ("  " if url else ""))
            rt.font.size = Pt(9)
            if url:
                ru = p.add_run(url)
                ru.font.size = Pt(9); ru.font.color.rgb = PB_LINK

    # ---------------------------------------------------------- assembly
    def build(self):
        m = self.meta
        if self.T.get("chrome", True):
            self.header_footer()
        block_files = sorted(glob.glob(os.path.join(self.dir, "blocks", "*.json")))
        if not block_files:
            raise FileNotFoundError(f"no blocks/*.json in {self.dir}")
        all_blocks = [compose.load_blocks(f) for f in block_files]
        section_titles = [b[1] for blocks in all_blocks for b in blocks
                          if b[0] in ("h1", "h2") and b[0] == ("h2" if self.T["cover"] == "meta_sheet" else "h1")]
        if self.T["cover"] == "meta_sheet":
            self.cover_meta_sheet(section_titles)
        else:
            self.cover()
        if m.get("toc", True) and self.T["cover"] != "meta_sheet":
            self.contents(section_titles)
        issues = []
        for f, blocks in zip(block_files, all_blocks):
            issues += qa.check_blocks(blocks, os.path.basename(f))
            cap_max = self.T.get("caption_max_words")
            if cap_max:
                for b in blocks:
                    if b[0] == "fig" and len(b[2].split()) > cap_max:
                        issues.append({"severity": "WARN", "rule": "caption-length",
                                       "where": os.path.basename(f),
                                       "detail": f"caption {len(b[2].split())} words > {cap_max} (template rule): {b[2][:60]}..."})
            self.render_blocks(blocks)
        self.references_section()
        if m.get("disclosure"):
            self.doc.add_heading("Disclosures and Method", level=1)
            for chunk in m["disclosure"].split("\n\n"):
                self.para(chunk, size=8.6, color=GREY,
                          align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)
        out = os.path.join(self.outdir, m["outname"] + ".docx")
        self.doc.save(out)
        issues += qa.check_docx(out)
        ok, fails, warns = qa.gate(issues)
        report = qa.format_report(issues)
        with open(os.path.join(self.dir, "qa_report.txt"), "w", encoding="utf-8") as fh:
            fh.write(report)
        if not ok:
            os.remove(out)
            raise SystemExit("QA FAIL - build blocked:\n" + report)
        return out, report

    def to_pdf(self, docx_path):
        if not shutil.which("soffice"):
            return None
        # Isolated profile per invocation: a stale LibreOffice profile lock
        # makes --convert-to exit 0 while writing nothing.
        profile = f"file:///tmp/lo_profile_{os.getpid()}"
        pdf_path = docx_path.replace(".docx", ".pdf")
        for _ in range(2):
            subprocess.run(["soffice", f"-env:UserInstallation={profile}",
                            "--headless", "--convert-to", "pdf",
                            "--outdir", self.outdir, docx_path],
                           check=True, capture_output=True, timeout=300)
            if os.path.exists(pdf_path):
                return pdf_path
        return None

    def measure_toc(self, pdf_path):
        """Measure section start pages from the rendered PDF; write toc_pages.json.
        Rebuild after this to print real page numbers in Contents.

        Uses poppler's pdftotext per page (the proven path from the original
        extract_toc.py); returns None when poppler is unavailable."""
        import re as _re
        if not shutil.which("pdftotext") or not shutil.which("pdfinfo"):
            return None
        info = subprocess.run(["pdfinfo", pdf_path], capture_output=True, text=True).stdout
        n = int(info.split("Pages:")[1].split()[0])
        raw, norm = {}, {}
        for p in range(1, n + 1):
            out = subprocess.run(["pdftotext", "-f", str(p), "-l", str(p),
                                  pdf_path, "-"], capture_output=True, text=True).stdout
            raw[p] = out
            norm[p] = _re.sub(r"\s+", " ", out)
        toc_page = next((p for p in range(2, n + 1)
                         if "Contents" in norm[p]), None)
        block_files = sorted(glob.glob(os.path.join(self.dir, "blocks", "*.json")))
        titles = [b[1] for f in block_files for b in compose.load_blocks(f) if b[0] == "h1"]
        pages = {}
        for t in titles:
            for p in range(2, n + 1):
                if p != toc_page and any(l.strip() == t for l in raw[p].splitlines()):
                    pages[t] = p
                    break
            else:
                for p in range(2, n + 1):
                    if p != toc_page and t in norm[p]:
                        pages[t] = p
                        break
        if not pages:
            return None
        with open(os.path.join(self.dir, "toc_pages.json"), "w", encoding="utf-8") as fh:
            json.dump(pages, fh, indent=2)
        return pages


def build_report(report_dir, pdf=True, measure=True):
    b = ReportBuilder(report_dir)
    docx_path, report = b.build()
    pdf_path = b.to_pdf(docx_path) if pdf else None
    if pdf_path and os.path.exists(pdf_path) and measure:
        pages = b.measure_toc(pdf_path)
        if pages:
            b2 = ReportBuilder(report_dir)
            docx_path, report = b2.build()
            pdf_path = b2.to_pdf(docx_path)
    return {"docx": docx_path, "pdf": pdf_path, "qa": report}
