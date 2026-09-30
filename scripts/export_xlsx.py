#!/usr/bin/env python3
"""
CONTEXT: ANTHROPIC INTELLIGENCE DESK — spreadsheet export of the record, in date order.
Writes data/anthropic_record.xlsx: Timeline (date, headline, one-sentence summary, source URL, tier, confidence, extracted fields),
plus one sheet per desk table and the conflict / open-item ledgers. Run: python3 scripts/export_xlsx.py
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D = lambda n: json.load(open(os.path.join(ROOT, "data", n)))
E, M, ENT, L = D("events.json"), D("metrics.json"), D("entities.json"), D("ledger.json")
HDR = PatternFill("solid", fgColor="15181E"); HF = Font(bold=True, color="FFFFFF", name="Arial", size=10); BF = Font(name="Arial", size=10); LF = Font(name="Arial", size=10, color="0563C1", underline="single")
TIER = {"T1": "C6EFCE", "T2": "DDEBF7", "T3": "FFF2CC", "T4": "F8CBAD", "T5": "F8CBAD"}

def sheet(wb, name, headers, rows, widths=None, url_cols=(), tier_col=None, wrap_cols=()):
    ws = wb.create_sheet(name[:31])
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HDR, HF, Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([("" if v is None else (json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)) for v in r])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BF; c.alignment = Alignment(vertical="top", wrap_text=(c.column - 1) in wrap_cols)
        for i in url_cols:
            c = row[i]
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value; c.font = LF
        if tier_col is not None:
            t = row[tier_col].value
            if t in TIER: row[tier_col].fill = PatternFill("solid", fgColor=TIER[t])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers):
        w = (widths or {}).get(i, min(max(len(str(h)) + 2, 12), 40))
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.row_dimensions[1].height = 30
    return ws

wb = Workbook(); ws0 = wb.active; ws0.title = "README"
for line in [
    "Anthropic Intelligence Desk: record as of " + L.get("generated", ""),
    "Timeline: every deduped event in date order, with a one-sentence summary, the source URL actually opened, the source tier and the extracted fields.",
    "Tiers: T1 primary document; T2 Bloomberg/Reuters/WSJ/FT/CNBC/official/PitchBook field; T3 single-outlet scoop or estimate; T4 aggregator; T5 social.",
    "Confidence: HIGH, MEDIUM, LOW, VERIFY (open check), DISPUTED (frozen conflict; both values shown, neither chosen).",
    "Basis: Anthropic reports revenue gross of cloud-partner resale; run-rate is not recognized revenue; 'up to' is not contracted; debt is not equity raised.",
    "PitchBook TTM 4Q2026 / 4Q2027 revenue fields are forward projections, not current revenue.",
    f"Events: {len(E)} merged from {L.get('raw_events')} raw findings across {len(D('agent_log.json')['agents'])} agents.",
]: ws0.append([line])
ws0.column_dimensions["A"].width = 150
for c in ws0["A"]: c.font = BF
ws0["A1"].font = Font(bold=True, name="Arial", size=12)

tl_headers = ["Date", "Headline", "One-sentence summary", "Source URL", "Publisher", "Tier", "Confidence", "Categories", "Independent sources", "Amount $M", "Post-money $M", "Pre-money $M",
              "Run-rate $M", "Period revenue $M", "Period", "Margin %", "Margin type", "Compute partner", "Chips", "GW", "Compute $M", "Compute term", "Product", "Price in $/Mtok", "Price out $/Mtok",
              "Governance", "Moat", "Strategy", "People", "Entities", "Headcount", "Customers", "Corroborating URLs", "Notes", "Found by", "Event id"]
rows = []
for e in sorted(E, key=lambda x: (x["date"], x["headline"])):
    x = e["extracted"]; cp = x.get("compute") or {}; pr = x.get("product") or {}
    rows.append([e["date"], e["headline"], e["summary"], e["source"].get("url"), e["source"].get("publisher"), e["source"].get("tier"), e["confidence"], "; ".join(e["category"]), e.get("independent_sources"),
                 x.get("amount_usd_m"), x.get("valuation_post_usd_m"), x.get("valuation_pre_usd_m"), x.get("revenue_run_rate_usd_m"), x.get("revenue_period_usd_m"), x.get("revenue_period"), x.get("margin_pct"), x.get("margin_type"),
                 cp.get("partner"), cp.get("chips"), cp.get("gw"), cp.get("usd_m"), cp.get("term"), pr.get("name"), pr.get("price_in_per_mtok"), pr.get("price_out_per_mtok"),
                 x.get("governance"), x.get("moat"), x.get("strategy"),
                 "; ".join(f"{p.get('name')} ({p.get('role')}, {p.get('move')})" for p in x.get("people", []) if isinstance(p, dict)),
                 "; ".join(f"{p.get('name')}{' [' + p.get('ticker') + ']' if p.get('ticker') else ''} ({p.get('type')})" for p in x.get("entities", []) if isinstance(p, dict)),
                 x.get("headcount"), x.get("customers"), " | ".join(c.get("url", "") for c in e.get("corroboration", [])), e.get("notes"), ", ".join(e.get("agents", [e["agent"]])), e["id"]])
sheet(wb, "Timeline", tl_headers, rows, widths={0: 11, 1: 44, 2: 70, 3: 34, 4: 18, 5: 6, 6: 11, 7: 22, 25: 40, 26: 40, 27: 40, 28: 30, 29: 40, 32: 40, 33: 50, 35: 40}, url_cols=(3,), tier_col=5, wrap_cols=(1, 2))

def gen(name, key, cols, widths=None, urls=(), tier=None):
    rs = M.get(key) or []
    if not rs: return
    sheet(wb, name, [c[0] for c in cols], [[r.get(c[1]) for c in cols] for r in rs], widths=widths, url_cols=urls, tier_col=tier)

gen("Rounds", "rounds", [("Round", "round"), ("Date", "date"), ("Size $M", "size_usd_m"), ("Pre $M", "pre_usd_m"), ("Post $M", "post_usd_m"), ("Leads", "leads"), ("Participants", "participants"), ("Price/share", "share_price_usd"), ("Status", "status"), ("PitchBook deal", "pitchbook_deal_id"), ("Source", "source_url"), ("Tier", "tier"), ("Confidence", "confidence")], widths={5: 40, 6: 50, 10: 40}, urls=(10,), tier=11)
gen("Run-rate", "run_rate_desk", [("Date", "date"), ("Run-rate $M", "run_rate_usd_m"), ("Basis", "basis"), ("Publisher", "publisher"), ("Source", "source_url"), ("Tier", "tier"), ("Confidence", "confidence")], widths={4: 50}, urls=(4,), tier=5)
gen("Period revenue", "period_revenue", [("Period", "period"), ("$M", "usd_m"), ("Kind", "kind"), ("Basis", "basis"), ("Source", "source_url"), ("Tier", "tier"), ("Confidence", "confidence")], widths={3: 40, 4: 50}, urls=(4,), tier=5)
gen("Margins & losses", "margins", [("Metric", "metric"), ("Period", "period"), ("Value", "value"), ("Basis", "basis"), ("Source", "source_url"), ("Tier", "tier"), ("Confidence", "confidence")], widths={2: 30, 3: 40, 4: 50}, urls=(4,), tier=5)
gen("Customers", "customers", [("Metric", "metric"), ("Value", "value"), ("Date", "date"), ("Source", "source_url"), ("Tier", "tier"), ("Confidence", "confidence")], widths={0: 36, 3: 50}, urls=(3,), tier=4)
gen("Mix", "mix", [("Segment", "segment"), ("Share/value", "share_or_value"), ("Date", "date"), ("Source", "source_url"), ("Tier", "tier"), ("Confidence", "confidence")], widths={1: 36, 3: 50}, urls=(3,), tier=4)
gen("Projections", "projections", [("Period", "period"), ("Metric", "metric"), ("Value", "value"), ("Who", "who"), ("Source", "source_url"), ("Tier", "tier"), ("Confidence", "confidence")], widths={4: 50}, urls=(4,), tier=5)
gen("Comparables", "comparables", [("Company", "company"), ("Metric", "metric"), ("Value", "value"), ("Basis", "basis"), ("Date", "date"), ("Source", "source_url"), ("Tier", "tier")], widths={3: 40, 5: 50}, urls=(5,), tier=6)
gen("Compute", "compute", [("Partner", "partner"), ("Chips", "chips"), ("$M", "usd_m"), ("$ basis", "usd_basis"), ("GW", "gw"), ("GW basis", "gw_basis"), ("Term", "term"), ("Status", "status"), ("Financing", "financing"), ("Source", "source_url")], widths={1: 28, 6: 18, 9: 50}, urls=(9,))
gen("Compute costs", "compute_costs", [("Metric", "metric"), ("Value", "value"), ("Basis", "basis"), ("Source", "source_url"), ("Tier", "tier")], widths={0: 36, 1: 30, 2: 40, 3: 50}, urls=(3,), tier=4)
gen("Pricing", "pricing", [("Model", "model"), ("Release", "release_date"), ("In $/Mtok", "price_in_per_mtok"), ("Out $/Mtok", "price_out_per_mtok"), ("Context", "context_window"), ("Class", "tier_class"), ("Later changes", "later_changes"), ("Source", "source_url"), ("Tier", "tier")], widths={0: 28, 6: 60, 7: 50}, urls=(7,), tier=8)
gen("Plans", "plans", [("Plan", "plan"), ("$/month", "price_usd_month"), ("Launched", "launched"), ("Changes", "changes"), ("Source", "source_url")], widths={3: 50, 4: 50}, urls=(4,))
gen("Board", "board", [("Name", "name"), ("Role", "role"), ("Appointed by", "appointed_by"), ("Joined", "joined"), ("Left", "left"), ("Source", "source_url"), ("Tier", "tier")], widths={0: 30, 5: 50}, urls=(5,), tier=6)
gen("People moves", "people_moves", [("Name", "name"), ("Role", "role"), ("Move", "move"), ("Date", "date"), ("From/to", "from_to"), ("Source", "source_url"), ("Tier", "tier")], widths={0: 28, 1: 34, 4: 30, 5: 50}, urls=(5,), tier=6)
gen("Offices", "offices", [("City", "city"), ("Opened", "opened"), ("Note", "note"), ("Source", "source_url")], widths={2: 50, 3: 50}, urls=(3,))
gen("RSP versions", "rsp", [("Version", "version"), ("Date", "date"), ("Key change", "key_change"), ("Source", "source_url")], widths={2: 70, 3: 50}, urls=(3,))
gen("Legal cases", "cases", [("Matter", "matter"), ("Type", "type"), ("Court/agency", "court_or_agency"), ("Docket", "docket"), ("Filed", "filed"), ("Role", "role"), ("Claims", "claims"), ("Key rulings", "key_rulings"), ("Amount $M", "amount_usd_m"), ("Status (Sep 2026)", "status_2026_09"), ("IPO risk note", "ipo_risk_note"), ("Sources", "source_urls"), ("Tier", "tier")], widths={0: 40, 6: 50, 7: 60, 9: 40, 10: 50, 11: 50}, tier=12)
gen("Prospectus facts", "prospectus_facts", [("Metric", "metric"), ("Value", "value"), ("Period/basis", "period_or_basis"), ("Reported by", "reported_by"), ("Published", "published"), ("Source", "source_url"), ("Tier", "tier"), ("Confidence", "confidence"), ("Note", "note")], widths={0: 36, 1: 36, 2: 30, 5: 50, 8: 50}, urls=(5,), tier=6)
gen("Market share", "market_share", [("Metric", "metric"), ("Value", "value"), ("Date", "date"), ("Publisher", "publisher"), ("Source", "source_url"), ("Tier", "tier")], widths={0: 40, 1: 30, 4: 50}, urls=(4,), tier=5)

sheet(wb, "Entities", ["Name", "Types", "Public", "Ticker", "Amount $M", "First date", "Linked events", "Notes", "Source"],
      [[r["name"], ", ".join(r["types"]), "yes" if r.get("public") else ("no" if r.get("public") is False else ""), r.get("ticker"), r.get("amount_usd_m"), r.get("first_date"), r.get("degree"), " | ".join(r.get("notes", [])), (r.get("sources") or [""])[0]] for r in ENT],
      widths={0: 34, 1: 26, 7: 70, 8: 50}, url_cols=(8,))
conf = (L.get("cross_agent_conflicts") or []) + (L.get("agent_conflicts") or [])
sheet(wb, "Conflicts (frozen)", ["Topic", "Value A", "Source A", "Value B", "Source B", "Note", "Raised by"], [[c.get("topic"), c.get("value_a"), c.get("source_a"), c.get("value_b"), c.get("source_b"), c.get("note"), c.get("agent", "merge")] for c in conf],
      widths={0: 50, 1: 40, 2: 40, 3: 40, 4: 40, 5: 60}, url_cols=(2, 4), wrap_cols=(0, 1, 3, 5))
sheet(wb, "Open items", ["Claim", "Why unverified", "Best lead", "Raised by"], [[o.get("claim"), o.get("why_unverified"), o.get("best_lead"), o.get("agent")] for o in (L.get("open_items") or [])], widths={0: 60, 1: 60, 2: 50}, wrap_cols=(0, 1))
if L.get("claims"):
    sheet(wb, "Validation ledger", ["Claim", "Draft value", "Re-derived value", "Verdict", "Tier", "Independent origins", "Confidence", "Sources", "Red-team note"],
          [[c.get("claim"), c.get("draft"), c.get("rederived"), c.get("verdict"), c.get("tier"), c.get("cross_check"), c.get("confidence"), c.get("sources"), c.get("red_team")] for c in L["claims"]],
          widths={0: 60, 1: 34, 2: 34, 7: 60, 8: 60}, tier_col=4, wrap_cols=(0, 1, 2, 7, 8))
out = os.path.join(ROOT, "data", "anthropic_record.xlsx"); wb.save(out)
print("wrote", out, "sheets:", ", ".join(wb.sheetnames))
