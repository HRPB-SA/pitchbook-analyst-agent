"""Assemble the Databricks initiation note (.docx).
US Letter, Georgia body, navy headings, pageBreakBefore H1s, dual-width DXA
tables, embedded 300-DPI charts, page-x-of-y footer with COI/embargo line.
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import content_a, content_b, content_c, content_d

HERE = os.path.dirname(os.path.abspath(__file__))
CHARTS = os.path.join(HERE, "..", "charts")
OUTDIR = os.path.join(HERE, "..", "output")
os.makedirs(OUTDIR, exist_ok=True)

NAVY = RGBColor(0x1F, 0x2A, 0x44)
SLATE = RGBColor(0x35, 0x50, 0x6E)
GREEN = RGBColor(0x1C, 0x5D, 0x46)
GREY = RGBColor(0x6B, 0x72, 0x80)
INK = RGBColor(0x26, 0x2B, 0x35)
SOFT_HEX = "F2F5F9"
NAVY_HEX = "1F2A44"
GREEN_HEX = "1C5D46"
BORDER_HEX = "C9D2DE"

CONTENT_W = 9360  # dxa (6.5in)

doc = Document()

# ---------------------------------------------------------------- page setup
sec = doc.sections[0]
sec.page_width, sec.page_height = Emu(12240 * 635), Emu(15840 * 635)
for a in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, a, Inches(1))
sec.different_first_page_header_footer = True

# ------------------------------------------------------------------- styles
st = doc.styles["Normal"]
st.font.name = "Georgia"
st.font.size = Pt(10.5)
st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(7)
st.paragraph_format.line_spacing = 1.12

def _style_heading(name, size, color, bold=True, before=10, after=6,
                   page_break=False, italic=False):
    s = doc.styles[name]
    s.font.name = "Georgia"
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.italic = italic
    s.font.color.rgb = color
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.page_break_before = page_break
    s.paragraph_format.keep_with_next = True
    # ensure latin font applies (headings inherit +Headings otherwise)
    rpr = s.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi"):
        rf.set(qn(attr), "Georgia")

_style_heading("Heading 1", 16, NAVY, page_break=True, before=0, after=10)
_style_heading("Heading 2", 12.5, NAVY, before=14, after=5)
_style_heading("Heading 3", 11, SLATE, before=10, after=4, italic=True)

# --------------------------------------------------------------- xml helpers
def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)

def set_cell_width(cell, dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW"); tcPr.append(tcW)
    tcW.set(qn("w:w"), str(dxa)); tcW.set(qn("w:type"), "dxa")

def set_table_width(table, dxa):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), str(dxa)); tblW.set(qn("w:type"), "dxa")

def table_borders(table, hexcolor=BORDER_HEX, sz="4"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), sz)
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), hexcolor)
        borders.append(e)
    tblPr.append(borders)

def row_no_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)
    if header:
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)

def cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom),
                      ("right", right)):
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)

def cell_text(cell, text, size=8.6, bold=False, color=INK, align="L",
              italic=False, font="Georgia"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"L": WD_ALIGN_PARAGRAPH.LEFT, "R": WD_ALIGN_PARAGRAPH.RIGHT,
                   "C": WD_ALIGN_PARAGRAPH.CENTER}[align]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.name = font; r.font.size = Pt(size)
    r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color

def add_field(paragraph, instr):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"
    r.append(t); fld.append(r)
    paragraph._p.append(fld)

# ------------------------------------------------------------ block renderers
def add_para(text, size=10.5, color=INK, italic=False, bold=False,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=7):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.color.rgb = color
    r.font.italic = italic; r.font.bold = bold
    return p

def add_bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.28)
        r = p.add_run(it)
        r.font.size = Pt(10.5); r.font.color.rgb = INK

def col_widths(header, rows, align):
    n = len(header)
    weights = []
    for j in range(n):
        longest = max([len(str(header[j]))] +
                      [len(str(r[j])) for r in rows])
        weights.append(min(max(longest, 6), 60))
    total = sum(weights)
    return [max(int(CONTENT_W * w / total), 700) for w in weights]

def add_table_block(spec):
    header, rows = spec["header"], spec["rows"]
    align = spec.get("align", "L" * len(header))
    if spec.get("title"):
        tp = add_para(spec["title"], size=9.5, color=NAVY, bold=True,
                      align=WD_ALIGN_PARAGRAPH.LEFT, space_after=3)
        tp.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=len(rows) + 1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_width(t, CONTENT_W)
    table_borders(t)
    cell_margins(t)
    widths = spec.get("widths") or col_widths(header, rows, align)
    row_no_split(t.rows[0], header=True)
    for r_ in t.rows[1:]:
        row_no_split(r_)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        set_cell_width(c, widths[j])
        shade(c, NAVY_HEX)
        cell_text(c, h, size=8.6, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
                  align="L" if align[j] == "L" else align[j])
    for i, row in enumerate(rows):
        last = (i == len(rows) - 1)
        summary = last and any("total" in str(x).lower() or "Composite" in str(x)
                               or "Weighted" in str(x) for x in row)
        for j, val in enumerate(row):
            c = t.rows[i + 1].cells[j]
            set_cell_width(c, widths[j])
            if summary:
                shade(c, "E2E8F0")
            elif i % 2 == 1:
                shade(c, SOFT_HEX)
            cell_text(c, str(val), size=8.4, align=align[j],
                      bold=summary, color=INK)
    if spec.get("source"):
        add_para(spec["source"], size=7.8, color=GREY, italic=True,
                 align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10)
    else:
        add_para("", size=4, space_after=4)

def add_fig(fname, caption):
    path = os.path.join(CHARTS, fname)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(6.5))
    add_para(caption, size=7.8, color=GREY, italic=True,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_after=12)

def add_box(title, body):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(t, CONTENT_W)
    table_borders(t, hexcolor="8F3421", sz="8")
    cell_margins(t, top=110, bottom=110, left=160, right=160)
    c = t.rows[0].cells[0]
    set_cell_width(c, CONTENT_W)
    shade(c, "FBF5F3")
    c.text = ""
    p1 = c.paragraphs[0]
    p1.paragraph_format.space_after = Pt(4)
    r = p1.add_run(title)
    r.font.name = "Georgia"; r.font.size = Pt(9.5); r.font.bold = True
    r.font.color.rgb = RGBColor(0x8F, 0x34, 0x21)
    p2 = c.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r2 = p2.add_run(body)
    r2.font.name = "Georgia"; r2.font.size = Pt(9.5); r2.font.color.rgb = INK
    add_para("", size=4, space_after=6)

def render(blocks):
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            doc.add_heading(b[1], level=1)
        elif kind == "h2":
            doc.add_heading(b[1], level=2)
        elif kind == "h3":
            doc.add_heading(b[1], level=3)
        elif kind == "p":
            add_para(b[1])
        elif kind == "bullets":
            add_bullets(b[1])
        elif kind == "table":
            add_table_block(b[1])
        elif kind == "fig":
            add_fig(b[1], b[2])
        elif kind == "box":
            add_box(b[1], b[2])
        else:
            raise ValueError(kind)

# =========================================================== header / footer
hdr = sec.header
hp = hdr.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("DATABRICKS, INC.  |  INITIATION OF COVERAGE  |  JULY 10, 2026")
hr.font.name = "Georgia"; hr.font.size = Pt(7.5); hr.font.color.rgb = GREY
hr.font.bold = True
# thin rule under header
pPr = hp._p.get_or_add_pPr()
pbdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "4")
bottom.set(qn("w:space"), "4"); bottom.set(qn("w:color"), BORDER_HEX)
pbdr.append(bottom); pPr.append(pbdr)

ftr = sec.footer
fp = ftr.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
f1 = fp.add_run("COI: Databricks is held in the author's managed book (8% core allocation)  |  "
                "Quality-valuation coefficient embargoed; per-point spread is the cleared expression  |  "
                "Not investment advice")
f1.font.name = "Georgia"; f1.font.size = Pt(6.8); f1.font.color.rgb = GREY
f1.font.italic = True
fp2 = ftr.add_paragraph()
fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp2.add_run("Page ")
r.font.name = "Georgia"; r.font.size = Pt(7.5); r.font.color.rgb = GREY
add_field(fp2, "PAGE")
r = fp2.add_run(" of ")
r.font.name = "Georgia"; r.font.size = Pt(7.5); r.font.color.rgb = GREY
add_field(fp2, "NUMPAGES")
for fldr in fp2._p.findall(qn("w:fldSimple")):
    for rr in fldr.findall(qn("w:r")):
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts"); fonts.set(qn("w:ascii"), "Georgia")
        fonts.set(qn("w:hAnsi"), "Georgia")
        szel = OxmlElement("w:sz"); szel.set(qn("w:val"), "15")
        colel = OxmlElement("w:color"); colel.set(qn("w:val"), "6B7280")
        rpr.append(fonts); rpr.append(szel); rpr.append(colel)
        rr.insert(0, rpr)

# first page (cover): empty header, confidentiality footer
fp_ftr = sec.first_page_footer.paragraphs[0]
fp_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp_ftr.add_run("CONFIDENTIAL  |  For institutional recipients only  |  Not for redistribution")
r.font.name = "Georgia"; r.font.size = Pt(7.5); r.font.color.rgb = GREY
r.font.italic = True

# ================================================================== COVER
def cover():
    add_para("", size=10, space_after=30)
    add_para("INITIATION OF COVERAGE", size=10, color=GREY, bold=True,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Databricks, Inc.")
    r.font.name = "Georgia"; r.font.size = Pt(30); r.font.bold = True
    r.font.color.rgb = NAVY
    add_para("Where the Models End and the Business Begins",
             size=14, color=SLATE, italic=True,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
    add_para("The highest-quality business in the Frontier Five, at the lowest quality-adjusted price. "
             "The IPO is the event that forces the market to reconcile the two.",
             size=10.5, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=16)

    # rating box
    t = doc.add_table(rows=1, cols=3)
    set_table_width(t, CONTENT_W)
    table_borders(t, hexcolor=NAVY_HEX, sz="8")
    cell_margins(t, top=130, bottom=130, left=120, right=120)
    heads = [("AIBQ RATING", "8.81  /  ELITE", "Highest in the Frontier Five; only Elite score in the cohort"),
             ("RECOMMENDATION", "CORE HOLDING", "8% allocation; conflict disclosed on page 2"),
             ("QUALITY-ADJUSTED PRICE", "$15.2B / AIBQ POINT", "Cheapest in cohort vs $118B (Anthropic), $188B (OpenAI)")]
    for j, (a, bb, cc) in enumerate(heads):
        c = t.rows[0].cells[j]
        set_cell_width(c, CONTENT_W // 3)
        shade(c, SOFT_HEX if j != 0 else "EAF2EE")
        c.text = ""
        p1 = c.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(a); r1.font.name = "Georgia"; r1.font.size = Pt(7.5)
        r1.font.bold = True; r1.font.color.rgb = GREY
        p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(bb); r2.font.name = "Georgia"; r2.font.size = Pt(13)
        r2.font.bold = True
        r2.font.color.rgb = GREEN if j == 0 else NAVY
        p3 = c.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(0)
        r3 = p3.add_run(cc); r3.font.name = "Georgia"; r3.font.size = Pt(7.3)
        r3.font.color.rgb = GREY
    add_para("", size=6, space_after=8)

    add_table_block({
        "title": "Snapshot (all figures re-verified July 10, 2026)",
        "header": ["Metric", "Value", "As of / source", "Tier"],
        "rows": [
            ["Last completed mark", "$134B (held flat, two closes)", "Series L; Dec 2025 / Feb 9, 2026; company + PitchBook", "T2"],
            ["Rumored round (not adopted)", "$165-175B, in talks, unclosed", "The Information, Jun 9, 2026; PB note Jul 7, 2026", "T3"],
            ["Revenue run-rate", "$6.9B, +80% YoY, accelerating", "Company via CNBC, Jun 16, 2026", "T2"],
            ["Gross margin", "74% (from >80%), guided lower", "Company disclosures, Jun 16, 2026", "T2"],
            ["Free cash flow", "Positive (TTM and FY2025)", "Company, Feb 9, 2026", "T2"],
            ["Net revenue retention", ">140%", "Company, Feb 9, 2026", "T2"],
            ["Run-rate multiple", "19.4x ($134B / $6.9B)", "Derived", "calc"],
            ["Capital efficiency (equity-only)", "0.34x ($6.9B / ~$20.2B)", "PitchBook 21-deal record, Jul 2026", "T2"],
            ["Lifetime capital", "~$29.5B (~$20.2B equity + ~$9.3B debt)", "PitchBook entity 59199-40", "T2"],
            ["IPO status", "No S-1; 2026 ruled out; 2027 earliest", "EDGAR Jul 10, 2026 (T1); Bloomberg TV, Jun 4, 2026", "T1/T2"],
        ],
        "align": "LLLC",
        "source": "Author: Harrison Rolfes, Senior Research Director  |  July 10, 2026  |  Single-name institutional initiation note. "
                  "Prepared for institutional-investor and venture-capital recipients.",
    })

cover()
render(content_a.BLOCKS)
render(content_b.BLOCKS)
render(content_c.BLOCKS)
render(content_d.BLOCKS)

out = os.path.join(OUTDIR, "Databricks_Initiation_Note_Jul2026.docx")
doc.save(out)
print("saved", out)

# ------------------------------------------------------------- em-dash audit
import zipfile
z = zipfile.ZipFile(out)
bad = []
for name in z.namelist():
    if name.endswith(".xml"):
        data = z.read(name).decode("utf8", "ignore")
        if "—" in data:
            bad.append(name)
print("EMDASH-CHECK:", "FAIL " + str(bad) if bad else "PASS (zero U+2014)")
