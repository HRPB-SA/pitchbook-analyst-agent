# xlhelp.py: shared openpyxl helpers for Workbook A (entry-fee-v5.xlsx) and Workbook B (greenfield-entry-cost.xlsx).
# House conventions (model-spec.md section 0; tools/prompts/model-prompt.md):
#   blue font  = hard-coded input; black = formula; green = link to another sheet / output KPI;
#   yellow fill = AJ (analyst judgment) / HOLE / low-confidence; Arial throughout; $M unless stated.
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

FONT = "Arial"
BLUE = "0000FF"; BLACK = "000000"; GREEN = "008000"; RED = "FF0000"; GREY = "595959"; WHITE = "FFFFFF"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HDR = PatternFill("solid", fgColor="D9E1F2")
SECT = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="EDEDED")

NUM = '$#,##0;($#,##0);-'          # $M
NUM1 = '$#,##0.0;($#,##0.0);-'
NUM2 = '$#,##0.00;($#,##0.00);-'
CNT = '#,##0;(#,##0);-'            # counts, GW*1000 etc.
CNT1 = '#,##0.0;(#,##0.0);-'
CNT2 = '#,##0.00;(#,##0.00);-'
CNT3 = '#,##0.000;(#,##0.000);-'
PCT = '0.0%'; PCT2 = '0.00%'
MULT = '0.0x'; MULT2 = '0.00x'; MULT3 = '0.000x'
DEC1 = '0.0'; DEC2 = '0.00'; DEC3 = '0.000'
DATEF = 'yyyy-mm-dd'
GEN = 'General'

NAMES = set()      # defined names (pre-registered so formula colouring can detect them)
REG = {}           # key -> absolute reference "'Sheet'!$C$5"
YC = {2023: "D", 2024: "E", 2025: "F", 2026: "G", 2027: "H", 2028: "I", 2029: "J", 2030: "K"}
YEARS = list(range(2023, 2031))
FYEARS = list(range(2026, 2031))

thin = Side(style="thin", color="808080")
B_TB = Border(top=thin, bottom=thin)
B_T = Border(top=thin)


def A(sheet, col, row):
    return f"'{sheet}'!${col}${row}"


def split_ref(ref):
    m = re.match(r"([A-Z]+)(\d+)$", ref)
    return m.group(1), int(m.group(2))


def R(key):
    """Absolute reference for a registered key."""
    if key not in REG:
        raise KeyError(f"REG key missing: {key}")
    return REG[key]


def RY(key, year):
    return R(f"{key}|{year}")


def has_key(key):
    return key in REG


def ref_uses_link(value):
    if not (isinstance(value, str) and value.startswith("=")):
        return False
    if "!" in value:
        return True
    for n in NAMES:
        if re.search(r"(?<![A-Za-z0-9_])%s(?![A-Za-z0-9_])" % re.escape(n), value):
            return True
    return False


class Sh:
    """Thin wrapper around a worksheet with the house styling and the address registry."""

    def __init__(self, wb, name, widths=None, default_width=None):
        self.wb = wb
        self.ws = wb.create_sheet(name)
        self.name = name
        self.ws.sheet_view.showGridLines = False
        if default_width:
            for i in range(1, 30):
                self.ws.column_dimensions[get_column_letter(i)].width = default_width
        for c, w in (widths or {}).items():
            self.ws.column_dimensions[c].width = w

    def put(self, ref, value, kind="label", nf=None, bold=False, italic=False, fill=None, wrap=False,
            size=9, align=None, color=None, name=None, key=None, border=None):
        c = self.ws[ref]
        c.value = value
        col = {"label": BLACK, "input": BLUE, "formula": BLACK, "link": GREEN, "output": GREEN,
               "note": GREY, "header": BLACK, "aj": BLUE, "hole": BLUE, "flag": RED, "text": BLACK,
               "section": WHITE}[kind]
        if kind in ("formula", "output") and ref_uses_link(value):
            col = GREEN
        if kind in ("aj", "hole") and isinstance(value, str) and value.startswith("="):
            col = GREEN if ref_uses_link(value) else BLACK
        if color:
            col = color
        c.font = Font(name=FONT, size=size, bold=(bold or kind in ("output", "header", "section")),
                      italic=(italic or kind == "note"), color=col)
        if kind in ("aj", "hole"):
            c.fill = YELLOW
        if kind == "header":
            c.fill = HDR
        if kind == "section":
            c.fill = SECT
        if fill:
            c.fill = fill
        if nf:
            c.number_format = nf
        if wrap or align:
            c.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="top" if wrap else None)
        if border:
            c.border = border
        if key:
            colL, rowN = split_ref(ref)
            REG[key] = A(self.name, colL, rowN)
        if name:
            colL, rowN = split_ref(ref)
            self.wb.defined_names[name] = DefinedName(name, attr_text=A(self.name, colL, rowN))
            NAMES.add(name)
        return c

    def section(self, row, text, ncols=12):
        for i in range(1, ncols + 1):
            self.ws.cell(row=row, column=i).fill = SECT
        self.put(f"A{row}", text, kind="section", size=10)

    def header(self, row, labels, start_col=1):
        for i, lab in enumerate(labels):
            col = get_column_letter(start_col + i)
            self.put(f"{col}{row}", lab, kind="header", border=B_TB, wrap=True)

    def note(self, ref, text, wrap=False):
        self.put(ref, text, kind="note", wrap=wrap)

    def yearhdr(self, row, years=YEARS, numeric=False, label="", cols=None):
        """Year header linked to 00_Assumptions row 3 (text labels) or row 2 (numeric)."""
        if label:
            self.put(f"A{row}", label, kind="header", border=B_TB)
        for y in years:
            col = cols[y] if cols else YC[y]
            src = R(f"YR|{y}") if numeric else R(f"YL|{y}")
            self.put(f"{col}{row}", f"={src}", kind="link", bold=True, border=B_TB, align="center",
                     nf=(GEN if numeric else None))


def register_names(names):
    for n in names:
        NAMES.add(n)


def tsrow(sh, row, key, label, unit, src, cells, kind="formula", nf=NUM, years=YEARS, bold=False, cols=None, srckind="note"):
    """Write a label + unit + source + per-year cells. cells: dict year -> value/formula (None = leave blank)."""
    sh.put(f"A{row}", label, bold=bold)
    if unit is not None:
        sh.put(f"B{row}", unit, kind="note")
    if src is not None:
        sh.put(f"C{row}", src, kind=srckind)
    for y in years:
        v = cells.get(y) if isinstance(cells, dict) else cells
        if v is None:
            continue
        col = cols[y] if cols else YC[y]
        k = kind
        if isinstance(v, str) and v.startswith("=") and kind in ("input", "aj", "hole"):
            k = kind
        if isinstance(v, str) and v == "HOLE":
            sh.put(f"{col}{row}", "HOLE", kind="hole", align="center")
            continue
        sh.put(f"{col}{row}", v, kind=k, nf=nf, bold=bold, key=(f"{key}|{y}" if key else None))


def driver(sh, row, key, label, unit, tag, bear, base, bull, nf=NUM, name=None, note=None, hole=False):
    """Scalar AJ driver row on 00_Assumptions: A label, B unit, C tag, D/E/F Bear/Base/Bull (yellow), G live."""
    kind = "hole" if hole else "aj"
    sh.put(f"A{row}", label)
    sh.put(f"B{row}", unit, kind="note")
    sh.put(f"C{row}", tag, kind="note")
    for col, v in (("D", bear), ("E", base), ("F", bull)):
        sh.put(f"{col}{row}", v, kind=kind, nf=nf, key=f"{key}_{ {'D':'BEAR','E':'BASE','F':'BULL'}[col] }")
    sh.put(f"G{row}", f"=CHOOSE(SW_SCEN,D{row},E{row},F{row})", kind="formula", nf=nf, bold=True, key=key, name=name)
    if note:
        sh.put(f"H{row}", note, kind="note")


def driver_ts(sh, row, key, label, unit, tag, bear, base, bull, years, nf=NUM, note=None, hole=False):
    """Time-series AJ driver block: row = live (CHOOSE), row+1..row+3 = Bear/Base/Bull rows (yellow)."""
    kind = "hole" if hole else "aj"
    sh.put(f"A{row}", label, bold=True)
    sh.put(f"B{row}", unit, kind="note")
    sh.put(f"C{row}", tag, kind="note")
    for i, (lab, vals) in enumerate((("Bear", bear), ("Base", base), ("Bull", bull))):
        rr = row + 1 + i
        sh.put(f"A{rr}", f"   {lab}", kind="note")
        for y in years:
            v = vals.get(y)
            if v is None:
                continue
            sh.put(f"{YC[y]}{rr}", v, kind=kind, nf=nf, key=f"{key}_{lab.upper()}|{y}")
    for y in years:
        c = YC[y]
        sh.put(f"{c}{row}", f"=CHOOSE(SW_SCEN,{c}{row+1},{c}{row+2},{c}{row+3})", kind="formula", nf=nf, bold=True, key=f"{key}|{y}")
    if note:
        sh.put(f"H{row}", note, kind="note")
    return row + 5


def drow(sh, row, id_, item, value, unit, rowid, status, tier, asof, basis, decay="STABLE", used="",
         note="", nf=NUM, kind="input", name=None, yellow=False, key=None):
    """01_Data row: A ID, B item, C source string, D value, E unit, F as-of, G basis, H status, I tier,
    J decay, K ledger row, L used-by, M note. Registers REG[id_] (or key) to the value cell."""
    src = f"{rowid} · {status} · {tier} · {asof} · {basis}"
    sh.put(f"A{row}", id_, kind="note")
    sh.put(f"B{row}", item)
    sh.put(f"C{row}", src, kind="note")
    k = kind
    if isinstance(value, str) and value.startswith("="):
        k = "formula"
    if yellow:
        k = "hole" if kind == "hole" else "aj"
    sh.put(f"D{row}", value, kind=k, nf=nf, key=(key or id_), name=name)
    sh.put(f"E{row}", unit, kind="note")
    sh.put(f"F{row}", asof, kind="note")
    sh.put(f"G{row}", basis, kind="note")
    sh.put(f"H{row}", status, kind="note")
    sh.put(f"I{row}", tier, kind="note")
    sh.put(f"J{row}", decay, kind="note")
    sh.put(f"K{row}", rowid, kind="note")
    sh.put(f"L{row}", used, kind="note")
    if note:
        sh.put(f"M{row}", note, kind="note")
    return row + 1
