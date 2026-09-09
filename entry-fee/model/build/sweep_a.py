# sweep_a.py: scenario and switch sweep for Workbook A. For each switch state: set the selector, save a temp copy,
# recalculate with LibreOffice (recalc.py), read the cached values, and record errors, the E10 gap, the relative-multiple
# sign, composites, CE ratio and the check-table FAIL count. Writes sweep_a.json and prints a table.
import os, sys, json, shutil, subprocess, warnings, tempfile
warnings.filterwarnings("ignore")
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(os.path.dirname(HERE), "entry-fee-v5.xlsx")
RECALC = "/mnt/skills/public/xlsx/scripts/recalc.py"
TMP = os.path.join(HERE, "sweep_tmp")
os.makedirs(TMP, exist_ok=True)

# selector cells: (sheet, cell, branches)
SEL = {
    "SW_SCEN": ("00_Assumptions", "C4", [1, 2, 3]),
    "SW_HAIRCUT": ("12_Conflicts", "D5", [1, 2]),
    "SW_OAI_BASIS": ("12_Conflicts", "D6", [1, 2]),
    "SW_IPO_OAI": ("12_Conflicts", "D7", [1, 2]),
    "SW_A_FY25": ("12_Conflicts", "D8", [1, 2]),
    "SW_A_FY24_NI": ("12_Conflicts", "D9", [1, 2]),
    "SW_A_2026_LOSS": ("12_Conflicts", "D10", [1, 2]),
    "SW_A_HC": ("12_Conflicts", "D11", [1, 2]),
    "SW_TPU_RATE": ("12_Conflicts", "D12", [1, 2, 3, 4]),
    "SW_TOTAL_RAISED_VIEW": ("12_Conflicts", "D13", [1, 2]),
    "SW_15B_FACILITY": ("12_Conflicts", "D14", [1, 2]),
    "SW_SSI": ("12_Conflicts", "D15", [1, 2]),
    "SW_GW_2028": ("12_Conflicts", "D16", [1, 2]),
    "SW_GOOGLE_VALUE": ("12_Conflicts", "D17", [1, 2]),
}
DEFAULT = {"SW_SCEN": 2, "SW_HAIRCUT": 1, "SW_OAI_BASIS": 1, "SW_IPO_OAI": 2, "SW_A_FY25": 1, "SW_A_FY24_NI": 2, "SW_A_2026_LOSS": 2,
           "SW_A_HC": 1, "SW_TPU_RATE": 2, "SW_TOTAL_RAISED_VIEW": 1, "SW_15B_FACILITY": 2, "SW_SSI": 1, "SW_GW_2028": 1, "SW_GOOGLE_VALUE": 1}

ERR = {"#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NULL!"}


def run_state(name, state):
    wb = load_workbook(SRC)
    for k, v in state.items():
        sh, cell, _ = SEL[k]
        wb[sh][cell].value = v
    p = os.path.join(TMP, f"{name}.xlsx")
    wb.save(p)
    out = subprocess.run([sys.executable, RECALC, p, "180"], capture_output=True, text=True)
    try:
        js = json.loads(out.stdout.strip().split("\n")[-1] if out.stdout.strip().startswith("{") is False else out.stdout)
    except Exception:
        js = {"status": "parse-error", "raw": out.stdout[-300:], "err": out.stderr[-300:]}
    wb2 = load_workbook(p, data_only=True)
    nerr = 0
    for ws in wb2.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ERR:
                    nerr += 1
    g = lambda s, c: wb2[s][c].value
    # FAIL count: locate the count row on 08_Output (first row after the table whose A starts with "Count")
    failcount = None
    o = wb2["08_Output"]
    for r in range(45, 220):
        v = o[f"A{r}"].value
        if isinstance(v, str) and v.startswith("Count of FAIL"):
            failcount = o[f"D{r}"].value
            break
    rec = {
        "state": state, "recalc_status": js.get("status"), "recalc_errors": js.get("total_errors"), "cached_errors": nerr,
        "scenario": g("00_Assumptions", "D4"), "FY26A": g("02_Revenue", "M54"), "FY26O": g("02_Revenue", "M64"),
        "gap28_lo": g("09_Obligations", "I88"), "gap28_hi": g("09_Obligations", "I89"), "gap28_live_env": g("09_Obligations", "I92"),
        "google_leg28": g("09_Obligations", "I86"), "relmult_live": g("06_Valuation", "D14"), "relmult_sign": ("+" if (g("06_Valuation", "D14") or 0) > 0 else "-"),
        "mult_A": g("06_Valuation", "D8"), "mult_O": g("06_Valuation", "D11"), "CE_A": g("05_Cash_Fund", "E43"), "CE_O": g("05_Cash_Fund", "G43"), "CE_ratio": g("05_Cash_Fund", "H43"),
        "comp_A_live": g("08_Output", "B37"), "comp_O_live": g("08_Output", "C37"), "ppt_A": g("08_Output", "B38"), "ppt_O": g("08_Output", "C38"),
        "A_2026_result": g("05_Cash_Fund", "D66"), "fac15": g("05_Cash_Fund", "D71"), "TR_view": g("05_Cash_Fund", "D24"), "fy25A": g("02_Revenue", "E25"),
        "failcount": failcount,
    }
    return rec


def main():
    results = {}
    results["BASE_DEFAULT"] = run_state("BASE_DEFAULT", dict(DEFAULT))
    for k, (sh, cell, branches) in SEL.items():
        for b in branches:
            if b == DEFAULT[k]:
                continue
            st = dict(DEFAULT); st[k] = b
            results[f"{k}={b}"] = run_state(f"{k}_{b}", st)
    # joint combinations of interest
    combos = {
        "Bear+proxy16.3": {"SW_SCEN": 1, "SW_GOOGLE_VALUE": 2, "SW_TPU_RATE": 4},
        "Bull+proxy6.9": {"SW_SCEN": 3, "SW_GOOGLE_VALUE": 2, "SW_TPU_RATE": 1},
        "proxy+15GW+12.5": {"SW_GOOGLE_VALUE": 2, "SW_GW_2028": 2, "SW_TPU_RATE": 3},
        "27pct+GROSS": {"SW_HAIRCUT": 2, "SW_OAI_BASIS": 2},
        "Bear+all-alt": {"SW_SCEN": 1, "SW_HAIRCUT": 2, "SW_OAI_BASIS": 2, "SW_IPO_OAI": 1, "SW_A_FY25": 2, "SW_A_FY24_NI": 1, "SW_A_2026_LOSS": 1, "SW_A_HC": 2, "SW_TPU_RATE": 4, "SW_TOTAL_RAISED_VIEW": 2, "SW_15B_FACILITY": 1, "SW_SSI": 2, "SW_GW_2028": 2, "SW_GOOGLE_VALUE": 2},
        "Bull+all-alt": {"SW_SCEN": 3, "SW_HAIRCUT": 2, "SW_OAI_BASIS": 2, "SW_IPO_OAI": 1, "SW_A_FY25": 2, "SW_A_FY24_NI": 1, "SW_A_2026_LOSS": 1, "SW_A_HC": 2, "SW_TPU_RATE": 1, "SW_TOTAL_RAISED_VIEW": 2, "SW_15B_FACILITY": 1, "SW_SSI": 2, "SW_GW_2028": 2, "SW_GOOGLE_VALUE": 2},
    }
    for name, delta in combos.items():
        st = dict(DEFAULT); st.update(delta)
        results[name] = run_state(name, st)
    with open(os.path.join(HERE, "sweep_a.json"), "w") as f:
        json.dump(results, f, indent=1, default=str)
    print(f"{'state':<26} {'recalc':<8} {'errs':>4} {'fails':>5} {'scen':<5} {'FY26A':>8} {'gap28lo':>10} {'gap28hi':>10} {'gleg28':>9} {'rel':>7} {'sign'} {'CEratio':>8} {'compA':>6} {'compO':>6}")
    for k, r in results.items():
        f = lambda v, w: (f"{v:{w}.0f}" if isinstance(v, (int, float)) else f"{str(v):>{w}}")
        print(f"{k:<26} {str(r['recalc_status']):<8} {r['cached_errors']:>4} {str(r['failcount']):>5} {str(r['scenario']):<5} {f(r['FY26A'],8)} {f(r['gap28_lo'],10)} {f(r['gap28_hi'],10)} {f(r['google_leg28'],9)} "
              f"{(f'{r['relmult_live']:+.1%}' if isinstance(r['relmult_live'], (int, float)) else 'n/a'):>7} {r['relmult_sign']:>4} {(f'{r['CE_ratio']:.2f}x' if isinstance(r['CE_ratio'], (int, float)) else 'n/a'):>8} "
              f"{(f'{r['comp_A_live']:.3f}' if isinstance(r['comp_A_live'], (int, float)) else 'n/a'):>6} {(f'{r['comp_O_live']:.3f}' if isinstance(r['comp_O_live'], (int, float)) else 'n/a'):>6}")


if __name__ == "__main__":
    main()
