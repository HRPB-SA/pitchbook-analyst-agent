# build_a.py: builds Workbook A (entry-fee-v5.xlsx) with openpyxl from the Analyst's model spec.
# Run: python3 build_a.py  -> model/entry-fee-v5.xlsx (formulas only; recalculate with recalc.py, then requote.py)
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from xlhelp import Sh, REG
import tabs_a1, tabs_a2, tabs_a3

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "entry-fee-v5.xlsx")

TABS = ["00_Assumptions", "01_Data", "02_Revenue", "03_Costs", "04_PL", "05_Cash_Fund", "06_Valuation", "07_Sensitivity",
        "08_Output", "09_Obligations", "10_Financing", "11_AIBQ", "12_Conflicts", "13_OutsideView"]
WIDTHS = {
    "00_Assumptions": {"A": 58, "B": 12, "C": 44, "D": 13, "E": 13, "F": 13, "G": 13, "H": 13, "I": 13, "J": 13, "K": 13},
    "01_Data": {"A": 13, "B": 62, "C": 52, "D": 14, "E": 12, "F": 16, "G": 40, "H": 14, "I": 10, "J": 12, "K": 16, "L": 22, "M": 30},
    "02_Revenue": {"A": 62, "B": 10, "C": 40, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12, "J": 12, "K": 12, "L": 14, "M": 14, "N": 12, "O": 14},
    "03_Costs": {"A": 66, "B": 10, "C": 38, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12, "J": 12, "K": 12},
    "04_PL": {"A": 60, "B": 10, "C": 34, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12, "J": 12, "K": 12},
    "05_Cash_Fund": {"A": 66, "B": 12, "C": 30, "D": 14, "E": 14, "F": 14, "G": 44, "H": 12, "I": 30, "J": 14, "K": 20},
    "06_Valuation": {"A": 62, "B": 12, "C": 34, "D": 14, "E": 16, "F": 14, "G": 14, "H": 14},
    "07_Sensitivity": {"A": 40, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14, "G": 16, "H": 14, "I": 12, "J": 12},
    "08_Output": {"A": 62, "B": 10, "C": 22, "D": 34, "E": 22, "F": 34, "G": 26, "H": 30},
    "09_Obligations": {"A": 44, "B": 22, "C": 26, "D": 34, "E": 30, "F": 30, "G": 14, "H": 24, "I": 18, "J": 22, "K": 14},
    "10_Financing": {"A": 8, "B": 60, "C": 16, "D": 40, "E": 44, "F": 22, "G": 14, "H": 40},
    "11_AIBQ": {"A": 60, "B": 9, "C": 10, "D": 10, "E": 70, "F": 30, "G": 26, "H": 10, "I": 12, "J": 8, "K": 40, "L": 14, "M": 20},
    "12_Conflicts": {"A": 24, "B": 44, "C": 20, "D": 9, "E": 16, "F": 16, "G": 12, "H": 12, "I": 70, "J": 16, "K": 60},
    "13_OutsideView": {"A": 48, "B": 40, "C": 50, "D": 36, "E": 50, "F": 30, "G": 24, "H": 24},
}


def build():
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {}
    for t in TABS:
        sheets[t] = Sh(wb, t, WIDTHS.get(t))
    # dependency order: data -> assumptions -> switches -> revenue -> obligations -> costs -> PL -> cash -> AIBQ -> valuation -> sensitivity -> financing -> outside view -> output
    tabs_a1.build_01(sheets["01_Data"])
    tabs_a1.build_00(sheets["00_Assumptions"])
    tabs_a1.build_12(sheets["12_Conflicts"])
    tabs_a2.build_02(sheets["02_Revenue"])
    tabs_a2.build_09(sheets["09_Obligations"])
    tabs_a2.build_03(sheets["03_Costs"])
    tabs_a2.build_04(sheets["04_PL"])
    tabs_a3.build_05(sheets["05_Cash_Fund"])
    tabs_a3.build_11(sheets["11_AIBQ"])
    tabs_a3.build_06(sheets["06_Valuation"])
    tabs_a3.build_07(sheets["07_Sensitivity"])
    tabs_a3.build_10(sheets["10_Financing"])
    tabs_a3.build_13(sheets["13_OutsideView"])
    tabs_a3.build_08(sheets["08_Output"])
    # Excel-legal names: LB01..LB10 collide with cell references (column LB); rename to LB_01..LB_10 in every formula and name.
    import re
    from openpyxl.workbook.defined_name import DefinedName
    pat = re.compile(r"(?<![A-Za-z0-9_'!$])LB(\d\d)([ab]?)(?![A-Za-z0-9_(])")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = pat.sub(lambda m: f"LB_{m.group(1)}{m.group(2)}", c.value)
    for n in list(wb.defined_names.keys()):
        if pat.fullmatch(n):
            dn = wb.defined_names[n]; del wb.defined_names[n]
            nn = pat.sub(lambda m: f"LB_{m.group(1)}{m.group(2)}", n)
            wb.defined_names[nn] = DefinedName(nn, attr_text=dn.attr_text)
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = "1F3864" if ws.title in ("00_Assumptions", "01_Data", "12_Conflicts") else "4472C4"
    wb.save(OUT)
    n = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    print("saved", OUT, "formulas:", n, "names:", len(wb.defined_names))
    return OUT


if __name__ == "__main__":
    build()
