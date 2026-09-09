#!/usr/bin/env python3
# CONTEXT: $100B ENTRY FEE v5 — Gate 3 workbook checks (orchestrator), run after the Model agent delivers.
# (1) recalc via LibreOffice (recalc.py) → zero errors; (2) XML scan for embargoed correlation strings and dashes;
# (3) named ranges for switches and LB01-LB10 exist; (4) cached values of the load-bearing outputs vs the blind
# derivation; (5) colour discipline sample: inputs blue, formulas black; hard-coded numbers inside formula cells.
import json, re, subprocess, sys, zipfile, pathlib
import openpyxl

BASE = pathlib.Path("/home/user/pitchbook-analyst-agent/entry-fee")
FILES = [BASE / "model/entry-fee-v5.xlsx", BASE / "model/greenfield-entry-cost.xlsx"]
EXPECT = {  # blind derivation + Analyst check values (Base defaults)
    "LB01": 65000, "LB02": 40000, "LB03": 124254.0, "LB04": 181216.5, "LB08": 24100,
}
SWITCHES = ["SW_SCEN", "SW_HAIRCUT", "SW_OAI_BASIS", "SW_IPO_OAI", "SW_A_FY25", "SW_A_FY24_NI", "SW_A_2026_LOSS",
            "SW_A_HC", "SW_TPU_RATE", "SW_GW_2028", "SW_TOTAL_RAISED_VIEW", "SW_15B_FACILITY", "SW_SSI", "SW_GOOGLE_VALUE"]
BLUE = {"FF0000FF", "000000FF", "0000FF"}

def recalc(f):
    r = subprocess.run(["python3", "/mnt/skills/public/xlsx/scripts/recalc.py", str(f), "180"], capture_output=True, text=True, timeout=400)
    try:
        j = json.loads(r.stdout[r.stdout.index("{"):])
    except Exception:
        return {"status": "no-json", "raw": (r.stdout + r.stderr)[-400:]}
    return j

def xml_scan(f):
    z = zipfile.ZipFile(f); hits = {"correl": 0, "-0.99": 0, "dash": 0, "r=": 0}
    for n in z.namelist():
        if n.endswith(".xml"):
            t = z.read(n).decode("utf8", "ignore")
            hits["correl"] += len(re.findall(r"(?i)correl", t)); hits["-0.99"] += t.count("-0.99") + t.count("−0.99")
            hits["dash"] += t.count("—") + t.count("–"); hits["r="] += len(re.findall(r"\br\s?=\s?-?0\.\d", t))
    return hits

def names(wb):
    out = {}
    try:
        for n, d in wb.defined_names.items():
            out[n] = d.attr_text
    except AttributeError:
        for d in wb.defined_names.definedName:
            out[d.name] = d.attr_text
    return out

def cell_by_name(wb, ref):
    m = re.match(r"'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)", ref)
    if not m: return None
    return wb[m.group(1)][f"{m.group(2)}{m.group(3)}"].value

for f in FILES:
    print("=" * 100); print(f.name)
    if not f.exists(): print("  MISSING"); continue
    rc = recalc(f); print("  recalc:", {k: rc.get(k) for k in ("status", "total_formulas", "total_errors")}, rc.get("error_summary") or "")
    print("  xml scan (must be 0 for correl/-0.99/r=/dash):", xml_scan(f))
    wbf = openpyxl.load_workbook(f)                   # formulas
    wbv = openpyxl.load_workbook(f, data_only=True)   # cached values (after recalc)
    nm = names(wbf); print("  sheets:", wbf.sheetnames)
    missing = [s for s in SWITCHES if s not in nm] if "entry-fee" in f.name else [s for s in ["SW_GF", "SW_GF_BACKSTOP", "SW_GF_REGION", "SW_SSI"] if s not in nm]
    print("  named ranges:", len(nm), "· missing switches:", missing)
    lb = {k: v for k, v in nm.items() if re.match(r"LB\d\d", k)}
    print("  LB names:", sorted(lb))
    for k, ref in sorted(lb.items()):
        try:
            v = cell_by_name(wbv, ref)
        except Exception as e:
            v = f"ERR {e}"
        exp = EXPECT.get(k[:4]); flag = "" if exp is None else ("OK" if (isinstance(v, (int, float)) and abs(v - exp) < 1) else f"MISMATCH expected {exp}")
        print(f"    {k} = {v}  {flag}")
    # colour discipline + hardcodes in formulas
    blue_inputs = black_formula = blue_formula = nonblue_input = 0; hard = []
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None: continue
                col = (c.font.color.rgb if c.font and c.font.color and c.font.color.type == "rgb" else None)
                if isinstance(v, str) and v.startswith("="):
                    if col in BLUE: blue_formula += 1
                    else: black_formula += 1
                    if ws.title not in ("00_Assumptions", "01_Data") and re.search(r"(?<![A-Z$!:\d.])\d{3,}(?:\.\d+)?(?![\d:])", v.replace("1000", "").replace("12", "")):
                        hard.append(f"{ws.title}!{c.coordinate}: {v[:60]}")
                elif isinstance(v, (int, float)):
                    if col in BLUE: blue_inputs += 1
                    else: nonblue_input += 1
    print(f"  formulas: {black_formula} black, {blue_formula} blue(!) · numeric inputs: {blue_inputs} blue, {nonblue_input} not blue")
    print(f"  formula cells with 3+ digit literals outside 00/01 tabs: {len(hard)} (first 12):")
    for h in hard[:12]: print("     ", h)
print("=" * 100)
tie = BASE / "model/tie-out.md"
print("tie-out.md:", "present" if tie.exists() else "MISSING", tie.stat().st_size if tie.exists() else "")
